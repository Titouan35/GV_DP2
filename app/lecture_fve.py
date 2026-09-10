"""Lecture de la Fiche de Validation d'Emprises (FVE) produite par le BE.

Le bureau d'études fournit, en plus des pièces du dossier, une FVE PowerPoint
qui contient déjà une bonne partie de ce que l'outil faisait ressaisir à la
main : puissance crête, module, point bas, inclinaison, largeur d'ombrière, et
surtout le photomontage d'insertion paysagère (la pièce DP6).

Deux principes, hérités de la lecture du plan de masse (app/lecture_plan.py) :

1. On PROPOSE, on n'impose pas. Les valeurs lues ne remplissent que les champs
   VIDES, elles sont renvoyées à l'interface qui les surligne, et l'utilisateur
   les corrige librement. Une FVE est un document d'avant-projet dont les
   valeurs bougent : l'importer en silence figerait une version périmée dans un
   dossier réglementaire, exactement le mécanisme qui a produit les
   incohérences d'adresse du 01/09/2026.
2. On ne DEVINE jamais. La FVE parle d'un « taux d'autoconsommation de 95 % » :
   on pourrait en déduire une autoconsommation avec vente du surplus, et se
   tromper. La destination de l'électricité reste donc un choix explicite de
   l'utilisateur à l'étape 3.

La localisation et les parcelles ne sont volontairement PAS touchées : elles
fonctionnent bien par géocodage et cadastre (décision de Florent, 01/09/2026).

Le gabarit de FVE est stable (confirmé par Florent) : la lecture s'appuie sur
les libellés du modèle. Tout libellé absent est signalé, jamais deviné.
"""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

logger = logging.getLogger(__name__)

# Libellés des deux images (avant / après) de la planche d'insertion. La liaison
# libellé -> image se fait par proximité VERTICALE : sur la FVE de référence,
# « Image source » est à 0,86 pouce et sa photo à 0,72 ; « Insertion
# Paysagère » à 3,92 et sa photo à 3,92.
# Deux variantes du gabarit circulent : la fiche « MARKET » (EXEMPLE/, relevé
# le 10/09/2026) étiquette le photomontage « IP-3D ». Sans cet alias, seule la
# photo avant était lue et la DP6 n'était jamais proposée.
LIBELLES_SOURCE = ("image source",)
LIBELLES_INSERTION = ("insertion paysag", "ip-3d", "ip 3d", "ip3d")
# Un libellé d'image est court ; un paragraphe qui cite le mot n'en est pas un.
LONGUEUR_MAX_LIBELLE = 40

# En dessous, une image est un pictogramme de la mise en page, pas une photo.
TAILLE_MINI_PHOTO = 200_000

# Titre de planche du gabarit : « 03 INSERTION PAYSAGERE », « 02 IMPLANTATION ».
_RE_TITRE_SECTION = re.compile(r"^\d{1,2}\s")


def _nombre(texte: str) -> float | None:
    """« 157,78 » ou « 3.5 » -> float. None si ce n'est pas un nombre."""
    try:
        return float(texte.replace(",", ".").strip())
    except (ValueError, AttributeError):
        return None


# Chaque règle : (clé du projet, expression, conversion).
# Les libellés viennent du gabarit FVE ; l'espace avant « : » est optionnel car
# PowerPoint insère parfois une espace insécable.
REGLES = [
    ("puissance_kwc", r"Puissance\s+crête\s*:\s*([\d\s,.]+)\s*kWc", _nombre),
    ("garde_au_sol_m", r"Point\s+bas\s*:\s*([\d\s,.]+)\s*m\b", _nombre),
    ("pente_deg", r"Inclinaison\s*:\s*([\d\s,.]+)\s*°", _nombre),
    # « Nombre de panneaux en largeur : 7 (12,3m) » : la cote entre parenthèses
    # est la largeur de l'ombrière, c'est-à-dire sa profondeur au sens du Cerfa.
    ("largeur_m", r"Nombre\s+de\s+panneaux\s+en\s+largeur\s*:[^(]*\(\s*([\d\s,.]+)\s*m\s*\)", _nombre),
]


def _texte_planches(prs) -> list[str]:
    """Texte de chaque forme, planche par planche (une entrée par forme)."""
    morceaux = []
    for planche in prs.slides:
        for forme in planche.shapes:
            if forme.has_text_frame and forme.text_frame.text.strip():
                morceaux.append(forme.text_frame.text)
            if getattr(forme, "has_table", False):
                for ligne in forme.table.rows:
                    morceaux.append(" | ".join(c.text for c in ligne.cells))
    return morceaux


def _module(morceaux: list[str]) -> dict:
    """Puissance unitaire et modèle du module.

    Le bloc « Modules » et le bloc « Onduleurs » portent les MÊMES libellés
    (« Marque-modèle », « Puissance nominale »). On les distingue par l'unité :
    le module est en Wc, l'onduleur en kVA.
    """
    for texte in morceaux:
        if "Marque-modèle" not in texte or "Wc" not in texte:
            continue
        champs = {}
        m = re.search(r"Puissance\s+nominale\s*:\s*([\d\s,.]+)\s*Wc", texte)
        if m:
            champs["module_puissance_wc"] = _nombre(m.group(1))
        m = re.search(r"Marque-modèle\s*:\s*([^\n|]+)", texte)
        if m:
            champs["module_modele"] = m.group(1).strip()
        return champs
    return {}


def _images_insertion(prs) -> dict:
    """Photo du site et photomontage d'insertion, par proximité des libellés.

    On renvoie l'image ENTIÈRE, pas le recadrage de la FVE : PowerPoint n'y
    montre qu'un bandeau (28 % de la surface sur la fiche de référence), choisi
    pour la mise en page de la fiche. L'image complète cadre bien mieux la
    planche DP6 du dossier.
    """
    resultat: dict[str, bytes] = {}
    for planche in prs.slides:
        libelles: list[tuple[float, str]] = []
        photos: list[tuple[float, float, bytes]] = []
        for forme in _formes(planche.shapes):
            haut = _pouces(forme.top)
            if forme.has_text_frame:
                minuscule = forme.text_frame.text.strip().lower()
                # Le TITRE de la planche (« 03 INSERTION PAYSAGERE ») contient
                # le même mot que le libellé de l'image et se trouve en haut,
                # donc plus près de la photo « source » que de l'insertion :
                # sans ce filtre, les deux rôles pointaient la même image.
                if _RE_TITRE_SECTION.match(minuscule) or len(minuscule) >= LONGUEUR_MAX_LIBELLE:
                    continue
                milieu = haut + _pouces(forme.height) / 2
                if any(l in minuscule for l in LIBELLES_SOURCE):
                    libelles.append((milieu, "source"))
                elif any(l in minuscule for l in LIBELLES_INSERTION):
                    libelles.append((milieu, "insertion"))
            # Picture ET PlaceholderPicture : toute forme qui porte une image.
            if hasattr(forme, "image"):
                try:
                    blob = forme.image.blob
                except (ValueError, AttributeError):
                    continue          # image liée, non embarquée
                if len(blob) >= TAILLE_MINI_PHOTO:
                    photos.append((haut, haut + _pouces(forme.height), blob))
        # Chaque photo ne sert qu'une fois : sans cela, deux libellés proches
        # d'une même photo lui attribuaient les deux rôles (avant = après).
        paires = sorted(
            (_distance(milieu, photo), i, role)
            for milieu, role in libelles
            for i, photo in enumerate(photos)
        )
        prises: set[int] = set()
        for _, i, role in paires:
            if role in resultat or i in prises:
                continue
            resultat[role] = photos[i][2]
            prises.add(i)
    return resultat


def _formes(formes):
    """Formes de la planche, y compris celles rangées dans des groupes."""
    for forme in formes:
        if getattr(forme, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _formes(forme.shapes)
        else:
            yield forme


def _pouces(emu) -> float:
    return Emu(emu or 0).inches


def _distance(y: float, photo: tuple[float, float, bytes]) -> float:
    """Écart vertical entre un libellé et une photo : nul s'il est À CÔTÉ.

    Comparer au seul bord haut était fragile : sur la fiche MARKET, « IP-3D »
    est à mi-hauteur de sa photo, un pouce sous son bord haut.
    """
    haut, bas, _ = photo
    if haut <= y <= bas:
        return 0.0
    return min(abs(y - haut), abs(y - bas))


def lire_fve(chemin: Path) -> dict:
    """Lit une FVE .pptx. Renvoie {champs, images, trouve, manquant}.

    Ne lève jamais : un fichier illisible rend un résultat vide et tracé, comme
    la lecture du plan de masse.
    """
    vide = {"champs": {}, "images": {}, "trouve": [], "manquant": []}
    if Path(chemin).suffix.lower() != ".pptx":
        return vide
    try:
        from pptx import Presentation

        prs = Presentation(str(chemin))
        morceaux = _texte_planches(prs)
        texte = "\n".join(morceaux)

        champs: dict[str, float | str] = {}
        manquant: list[str] = []
        for cle, motif, convertir in REGLES:
            m = re.search(motif, texte, re.IGNORECASE)
            valeur = convertir(m.group(1)) if m else None
            if valeur is not None:
                champs[cle] = valeur
            else:
                manquant.append(cle)
        champs.update(_module(morceaux))

        return {
            "champs": champs,
            "images": _images_insertion(prs),
            "trouve": sorted(champs),
            "manquant": manquant,
        }
    except Exception:      # FVE corrompue, chiffrée, ou gabarit inattendu
        logger.warning("Lecture de la FVE impossible (%s)", Path(chemin).name,
                       exc_info=True)
        return vide


def appliquer_lecture(projet, lecture: dict) -> list[str]:
    """Remplit les champs VIDES de l'ombrière. Renvoie les champs proposés.

    Ne réécrit jamais une valeur déjà saisie : la FVE propose, l'utilisateur
    dispose. `module_modele` n'a pas de champ dédié dans le modèle, il alimente
    `type_module` qui sert d'information libre.
    """
    proposes: list[str] = []
    correspondances = {"module_modele": "type_module"}
    for cle, valeur in (lecture.get("champs") or {}).items():
        attribut = correspondances.get(cle, cle)
        if not hasattr(projet.ombriere, attribut):
            continue
        if getattr(projet.ombriere, attribut) in (None, ""):
            setattr(projet.ombriere, attribut, valeur)
            proposes.append(attribut)
    return proposes
