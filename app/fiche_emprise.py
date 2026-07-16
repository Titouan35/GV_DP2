"""Fiche de validation d'emprise (visuel commercial, hors dossier réglementaire).

Reprend la trame de `EXEMPLE/MARKET_FICHE_VALIDATION_EMPRISE_v0` : un slide
16:9 à la charte GV avec la photo du site (« avant ») et l'insertion retenue
(« après »), étiquetée « visuel IA ». Ce livrable NE fait PAS partie du dossier
DP : il sert aux fiches d'emprise et présentations client (PLAN §4, §6 bis).
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from . import config

NAVY = RGBColor(0x00, 0x24, 0x55)
VIOLET = RGBColor(0x77, 0x6D, 0xF8)
MUTED = RGBColor(0x64, 0x74, 0x8B)
AMBRE = RGBColor(0xB4, 0x54, 0x09)
POLICE = "Poppins"

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _texte(slide, x, y, w, h, contenu, taille=18, gras=False, couleur=NAVY, centre=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lignes = contenu.split("\n") if isinstance(contenu, str) else contenu
    for i, ligne in enumerate(lignes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ligne
        p.font.size = Pt(taille)
        p.font.bold = gras
        p.font.color.rgb = couleur
        p.font.name = POLICE
        if centre:
            p.alignment = 2
    return box


def _gradient_png(assets: Path) -> Path:
    from .planches.base import gradient_h
    chemin = assets / "_gradient.png"
    if not chemin.exists():
        gradient_h(1600, 24).save(chemin)
    return chemin


def _image_dans_zone(slide, chemin_img: Path, zone):
    """Colle l'image ajustée (contain) dans la zone (x, y, w, h) en EMU."""
    with Image.open(chemin_img) as im:
        iw, ih = im.size
    x, y, w, h = zone
    ratio = min(w / iw, h / ih)
    iw2, ih2 = int(iw * ratio), int(ih * ratio)
    cx = int(x + (w - iw2) / 2)
    cy = int(y + (h - ih2) / 2)
    slide.shapes.add_picture(str(chemin_img), cx, cy, iw2, ih2)


def generer_fiche(projet: dict) -> Path:
    """Construit la fiche d'emprise PPTX. Requiert l'image d'insertion retenue."""
    from .insertion_ia import image_kit  # local : évite un cycle d'import

    ins = projet.get("insertion") or {}
    retenue = ins.get("retenue")
    if not retenue:
        raise ValueError("Aucune image d'insertion retenue : importez et retenez une image (étape 5).")
    chemin_insertion = config.PROJETS_DIR / retenue
    if not chemin_insertion.exists():
        raise ValueError("Image retenue introuvable sur le disque.")

    projet_id = projet["id"]
    assets = config.assets_dir(projet_id)
    loc = projet.get("localisation") or {}

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # bandeau
    slide.shapes.add_picture(str(_gradient_png(assets)), 0, 0, SLIDE_W, Emu(60000))
    _texte(slide, Inches(0.5), Inches(0.22), Inches(9.5), Inches(0.6),
           "Fiche de validation d'emprise", taille=24, gras=True)
    sous = " · ".join(filter(None, [projet.get("nom"), loc.get("commune"),
                                     f"INSEE {loc.get('code_insee')}" if loc.get("code_insee") else None]))
    _texte(slide, Inches(0.5), Inches(0.78), Inches(11), Inches(0.4), sous, taille=12, couleur=MUTED)
    _texte(slide, Inches(11.2), Inches(0.28), Inches(1.9), Inches(0.5),
           "greenvolt next", taille=14, gras=True, couleur=VIOLET)

    # photo du site (avant) — optionnelle
    photo = image_kit(projet, "photo")
    haut = Inches(1.25)
    if photo:
        _texte(slide, Inches(0.5), haut, Inches(3), Inches(0.35), "Photo du site", taille=13, gras=True)
        _image_dans_zone(slide, photo, (Inches(0.5), Inches(1.62), Inches(12.33), Inches(2.55)))
        bas_titre = Inches(4.35)
        bas_zone = (Inches(0.5), Inches(4.72), Inches(12.33), Inches(2.35))
    else:
        bas_titre = Inches(1.25)
        bas_zone = (Inches(0.5), Inches(1.62), Inches(12.33), Inches(4.9))

    # insertion (après)
    _texte(slide, Inches(0.5), bas_titre, Inches(6), Inches(0.35),
           "Insertion projetée", taille=13, gras=True)
    _texte(slide, Inches(2.4), bas_titre, Inches(4), Inches(0.35),
           "visuel IA — usage commercial", taille=11, gras=True, couleur=AMBRE)
    _image_dans_zone(slide, chemin_insertion, bas_zone)

    # pied de page (garde-fou réglementaire)
    _texte(slide, Inches(0.5), Inches(7.12), Inches(12.3), Inches(0.32),
           f"Visuel généré par IA, à usage commercial — ne constitue pas la pièce DP6 du "
           f"dossier réglementaire (fournie par le BE). Édité le {date.today().strftime('%d/%m/%Y')} avec GV_DP.",
           taille=9, couleur=MUTED)

    chemin = assets / f"Fiche_emprise_{projet_id}.pptx"
    prs.save(str(chemin))
    return chemin
