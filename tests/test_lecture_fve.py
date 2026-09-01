"""Lecture de la Fiche de Validation d'Emprises du bureau d'études.

Le gabarit est stable (confirmé par Florent le 01/09/2026), la lecture s'appuie
donc sur ses libellés. Ces tests construisent une FVE minimale en mémoire : ils
ne dépendent d'aucun fichier client.
"""
from __future__ import annotations

import io
import random

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import lecture_fve
from app.models import Projet


def _image(graine: int, taille=(900, 700)) -> io.BytesIO:
    """Photo assez lourde pour dépasser TAILLE_MINI_PHOTO (200 Ko).

    Le bruit pseudo-aléatoire est indispensable : un dégradé se compresse à
    quelques kilo-octets et l'image serait prise pour un pictogramme de mise en
    page, donc ignorée par le lecteur.
    """
    alea = random.Random(graine)
    octets = bytes(alea.getrandbits(8) for _ in range(taille[0] * taille[1] * 3))
    buf = io.BytesIO()
    Image.frombytes("RGB", taille, octets).save(buf, "PNG")
    buf.seek(0)
    return buf


def fve_de_test(tmp_path, textes=None, avec_images=True):
    """Construit une FVE au format du gabarit BE."""
    prs = Presentation()
    vierge = prs.slide_layouts[6]

    planche = prs.slides.add_slide(vierge)
    defaut = [
        "Marque-modèle : AIKO-A490-MCES4Dw\nPuissance nominale : 490 Wc\nNombre : 322",
        "Marque-modèle : SUNGROW 125 CX-P2\nPuissance nominale : 125 kVA\nNombre : 1",
        "Point bas : 3,5 m\nInclinaison : 6°\nNombre de panneaux en largeur : 7 (12,3m)",
        "Puissance crête : 157,78 kWc\nPuissance AC : 125 kVA",
    ]
    for i, texte in enumerate(textes if textes is not None else defaut):
        zone = planche.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(0.8))
        zone.text_frame.text = texte

    if avec_images:
        insertion = prs.slides.add_slide(vierge)
        titre = insertion.shapes.add_textbox(Inches(0.6), Inches(0.58), Inches(4), Inches(0.4))
        titre.text_frame.text = "03 INSERTION PAYSAGERE"
        lab1 = insertion.shapes.add_textbox(Inches(4.3), Inches(0.86), Inches(2), Inches(0.3))
        lab1.text_frame.text = "Image source"
        insertion.shapes.add_picture(_image(10), Inches(4.85), Inches(0.72), Inches(7), Inches(2.8))
        lab2 = insertion.shapes.add_textbox(Inches(4.3), Inches(3.92), Inches(2), Inches(0.3))
        lab2.text_frame.text = "Insertion Paysagère"
        insertion.shapes.add_picture(_image(200), Inches(4.87), Inches(3.92), Inches(7), Inches(2.7))

    chemin = tmp_path / "FVE.pptx"
    prs.save(str(chemin))
    return chemin


# ------------------------------------------------------------------ lecture

def test_les_caracteristiques_sont_lues(tmp_path):
    lecture = lecture_fve.lire_fve(fve_de_test(tmp_path))
    assert lecture["champs"]["puissance_kwc"] == 157.78
    assert lecture["champs"]["garde_au_sol_m"] == 3.5
    assert lecture["champs"]["pente_deg"] == 6.0
    assert lecture["champs"]["largeur_m"] == 12.3
    assert lecture["manquant"] == []


def test_le_module_est_distingue_de_l_onduleur(tmp_path):
    """Les deux blocs portent les MÊMES libellés (« Marque-modèle »,
    « Puissance nominale ») : seule l'unité les sépare, Wc contre kVA. Sans
    cela, l'onduleur de 125 kVA devenait un module de 125 Wc."""
    champs = lecture_fve.lire_fve(fve_de_test(tmp_path))["champs"]
    assert champs["module_puissance_wc"] == 490.0
    assert champs["module_modele"] == "AIKO-A490-MCES4Dw"
    assert "SUNGROW" not in str(champs)


def test_les_deux_images_sont_distinctes(tmp_path):
    """Le TITRE de planche (« 03 INSERTION PAYSAGERE ») contient le mot du
    libellé et se trouve plus près de la photo source : sans filtre sur les
    titres de section, les deux rôles pointaient la même image."""
    images = lecture_fve.lire_fve(fve_de_test(tmp_path))["images"]
    assert set(images) == {"source", "insertion"}
    assert images["source"] != images["insertion"]


def test_une_fve_sans_planche_insertion_ne_plante_pas(tmp_path):
    lecture = lecture_fve.lire_fve(fve_de_test(tmp_path, avec_images=False))
    assert lecture["images"] == {}
    assert lecture["champs"]["puissance_kwc"] == 157.78


def test_un_gabarit_inconnu_signale_ce_qui_manque(tmp_path):
    """Si le BE change ses libellés, on ne devine pas : on dit ce qu'on n'a
    pas trouvé."""
    lecture = lecture_fve.lire_fve(fve_de_test(tmp_path, textes=["Rien d'exploitable ici."]))
    assert lecture["champs"] == {}
    assert set(lecture["manquant"]) == {"puissance_kwc", "garde_au_sol_m",
                                        "pente_deg", "largeur_m"}


def test_un_fichier_qui_n_est_pas_un_pptx_est_ignore(tmp_path):
    faux = tmp_path / "pas_une_fve.pdf"
    faux.write_bytes(b"%PDF-1.4")
    assert lecture_fve.lire_fve(faux)["champs"] == {}


def test_un_pptx_corrompu_ne_leve_pas(tmp_path):
    """Un fichier illisible rend une lecture vide et tracée, comme le plan."""
    casse = tmp_path / "casse.pptx"
    casse.write_bytes(b"PK\x03\x04" + b"\x00" * 100)
    assert lecture_fve.lire_fve(casse) == {"champs": {}, "images": {},
                                           "trouve": [], "manquant": []}


# --------------------------------------------------------------- application

def test_seuls_les_champs_vides_sont_remplis(tmp_path):
    """La FVE PROPOSE. Une valeur déjà saisie par l'utilisateur ne doit jamais
    être écrasée : c'est un document d'avant-projet, ses valeurs bougent."""
    # Arrange
    projet = Projet(nom="X")
    projet.ombriere.puissance_kwc = 200.0        # déjà saisi à la main

    # Act
    proposes = lecture_fve.appliquer_lecture(projet, lecture_fve.lire_fve(fve_de_test(tmp_path)))

    # Assert
    assert projet.ombriere.puissance_kwc == 200.0
    assert "puissance_kwc" not in proposes
    assert projet.ombriere.pente_deg == 6.0
    assert "pente_deg" in proposes


def test_la_destination_de_l_energie_n_est_jamais_deduite(tmp_path):
    """La FVE parle d'un « taux d'autoconsommation de 95 % ». On pourrait en
    déduire une vente du surplus, et se tromper : le projet de référence est en
    autoconsommation TOTALE. Ce choix reste explicite."""
    # Arrange
    textes = ["Puissance cible à 95% Taux AC [kWc] | 286 | 157,75",
              "Puissance crête : 157,78 kWc"]

    # Act
    projet = Projet(nom="X")
    lecture_fve.appliquer_lecture(projet, lecture_fve.lire_fve(fve_de_test(tmp_path, textes=textes)))

    # Assert
    assert projet.ombriere.destination_energie is None


def test_la_localisation_n_est_jamais_touchee(tmp_path):
    """Décision de Florent : le géocodage et le cadastre fonctionnent bien,
    la FVE ne doit pas s'en mêler."""
    # Arrange
    textes = ["Adresse\n68 Rue du Jura, 01460 Montréal-la-Cluse",
              "Puissance crête : 157,78 kWc"]
    projet = Projet(nom="X")

    # Act
    lecture_fve.appliquer_lecture(projet, lecture_fve.lire_fve(fve_de_test(tmp_path, textes=textes)))

    # Assert
    assert projet.localisation.adresse is None
    assert projet.localisation.parcelles == []
